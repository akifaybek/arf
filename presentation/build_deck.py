"""Arf sunumu. Yalnızca presentation/ içine yazar; proje modüllerini yüklemez.

python presentation/build_deck.py  (python-pptx gerekir)
"""
from pathlib import Path
import ast
import json
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent
BG, PANEL, PAPER = "071F23", "0C2B30", "E8EBDF"
FG, SAGE, MUTED = "EEF1E8", "B5D0BD", "ADC0BA"
INK, PAPER_MUTED, RULE = "16312E", "4D655D", "35524F"
GREEN, AMBER, RED = "2BB673", "B17C17", "BF4746"
SANS, SERIF, MONO = "Arial", "Georgia", "Courier New"
GITHUB = "https://github.com/akifaybek/arf"
DEMO = "https://arf-icra-demo.copperspoon2.chatgpt.site"
CONTRACT = "0xddE4C78F831f082E8AFDF716266bd1338B591159"
EXPLORER = "https://explorer-test.avax.network/c-chain/tx/"
# Sabit kanıt listesini AST ile okur; app.py import edilmez, ortam dosyası açılmaz.
source = ast.parse((ROOT.parent / "app.py").read_text())
proofs = next(ast.literal_eval(n.value) for n in source.body if isinstance(n, ast.Assign)
              and any(isinstance(t, ast.Name) and t.id == "ONCHAIN_PROOFS" for t in n.targets))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = "Arf — İşlem öncesi icra ve likidite kalkanı"
prs.core_properties.subject = "OKX Agent Trade Kit × Avalanche Fuji"
prs.core_properties.author = "Arf"
records = []


def block(s, x, y, w, h, color):
    obj = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = RGBColor.from_string(color)
    obj.line.fill.background()
    return obj


def line(s, x1, y1, x2, y2, color=RULE, width=1, arrow=False):
    obj = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.line.color.rgb = RGBColor.from_string(color)
    obj.line.width = Pt(width)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        obj.line._get_or_add_ln().append(end)
    return obj


def text(s, x, y, w, h, value, size=24, color=None, font=SANS, bold=False,
         align=PP_ALIGN.LEFT, link=None):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = obj.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = Inches(.02)
    for i, part in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        p.line_spacing = 1.05
        run = p.add_run()
        run.text = part
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color or FG)
        if link:
            run.hyperlink.address = link
    return obj


def slide(title, note, seconds, light=False, heading=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg, fg, muted = (PAPER, INK, PAPER_MUTED) if light else (BG, FG, MUTED)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    if heading:
        text(s, .78, .67, 11.8, 1.12, title, 40, fg, SERIF)
    line(s, .78, 6.94, 12.55, 6.94, PAPER_MUTED if light else RULE, .7)
    text(s, .8, 7.08, 10.9, .25, "Arf   /   OKX Agent Trade Kit × Avalanche   /   17 Ekim 2026", 11, muted)
    text(s, 12.04, 7.05, .48, .3, f"{len(prs.slides):02}", 12, muted, MONO, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = note
    records.append({"slide":len(prs.slides),"title":title,"target_seconds":seconds,
                    "notes":note,"note_characters":len(note),"note_words":len(note.split())})
    return s


# 1 / Kimlik
s = slide("Arf", "Arf, AI trading agent'lar için işlem öncesi icra ve likidite kalkanı; ATK verisini Avalanche üzerindeki bir devre kesiciye bağlıyorum. Adı Cahit Arf ve Arf değişmezine gönderme yapıyor, çünkü sistem tek bir değişmezi koruyor: emir tutarı imzalı güvenli tutarı aşamaz. Bu sunumda problemi, kurduğum akışı ve kötü senaryoda ne olduğunu göstereceğim; canlı uygulama yerelde, internetteki paylaşılabilir sürüm ise açıkça etiketlenmiş kayıtlı bir simülasyon.", 30, heading=False)
block(s, 9.15, 0, 4.183333, 6.93, PAPER)
text(s, .75, 1.05, 7.8, 1.65, "Arf", 112, FG, SERIF)
text(s, .83, 3.06, 7.65, 1.63, "İşlem öncesi icra\nve likidite kalkanı", 36, FG)
text(s, .85, 5.36, 7.4, .7, "AI trading agent’lar için", 23, MUTED)
text(s, 9.53, 1.8, 3.4, .7, "emir tutarı", 25, INK, MONO, align=PP_ALIGN.CENTER)
text(s, 9.53, 2.66, 3.4, 1.5, "≤", 100, INK, SERIF, align=PP_ALIGN.CENTER)
text(s, 9.43, 4.37, 3.6, .7, "maxSafeSize", 25, INK, MONO, align=PP_ALIGN.CENTER)
text(s, .85, 6.3, 7.3, .32, "Cahit Arf’a ve değişmez fikrine gönderme.", 16, MUTED)

# 2 / Problem
s = slide("Aynı tahta. Çok farklı icra maliyeti.", "ATK ile binlerce trading agent yazılacak; tahtanın emri kaldırıp kaldıramayacağını bilmeyen agent kendini batırır. Hackathon ölçüm notundaki BIGTIME örneğinde bin USDT için kayma on virgül sekiz baz puanken on bin USDT için bin iki yüz seksen dört baz puan; tutar on kat, kayma oranı yüz kattan fazla büyüyor. Bunlar 17 Ekim tarihli tek bir tahta örneği, canlı fiyat veya genel performans garantisi değil; bu nedenle kontrolün her emirden önce tekrarlanması gerekiyor.", 35, light=True)
text(s, .85, 2.08, 5.45, .55, "BIGTIME-USDT", 22, PAPER_MUTED)
line(s, 6.55, 2.17, 6.55, 5.38, "B7C5BA")
text(s, .85, 2.88, 5.1, .55, "1.000 USDT", 24, INK)
text(s, .8, 3.55, 5.5, 1.15, "10,8", 76, INK, MONO)
text(s, .88, 4.92, 5.1, .5, "bps kayma", 23, PAPER_MUTED)
text(s, 7.15, 2.88, 5.1, .55, "10.000 USDT", 24, INK)
text(s, 7.1, 3.55, 5.2, 1.15, "1284", 76, INK, MONO)
text(s, 7.18, 4.92, 5.1, .5, "bps kayma", 23, PAPER_MUTED)
text(s, .85, 5.9, 11.5, .57, "10× tutar, 100× üzerinde kayma oranı", 28, INK)
text(s, .86, 6.57, 11.5, .24, "ÖLÇÜLEN: hackathon ölçüm notu. Canlı tahta değişir.", 14, PAPER_MUTED)

# 3 / Çözüm
s = slide("Bir emir için tek değişmez.", "Arf’ın sorduğu soru basit: bu emir bu tahtada kaça patlar ve en fazla ne kadarlık emir güvenli? check_execution sembolü, yönü ve tutarı alır; deterministik risk skorunu, güvenli tutarı ve gerekiyorsa bölme önerisini üretir, taze sonuç oracle tarafından imzalanır ve kontrat bu sınırı uygular. Buradaki güvenli tutar anlık tahta ölçümüdür, gelecekteki dolum garantisi değildir ve kontrat OKX üzerinde kendisi emir yürütmez.", 30)
text(s, .85, 2.0, 11.5, .55, "Her emirden önce: ölçüm, imza ve zincir kontrolü.", 27, MUTED)
block(s, .8, 3.05, 11.75, 1.68, PAPER)
text(s, 1.08, 3.48, 11.15, .85, "emir tutarı ≤ maxSafeSize", 44, INK, MONO, align=PP_ALIGN.CENTER)
text(s, .86, 5.45, 11.55, .7, "Tahtanın taşıyamadığı emir, sınırda durur.", 30, FG, SERIF)
text(s, .87, 6.31, 11.5, .35, "check_execution(symbol, side, size)", 20, SAGE, MONO)

# 4 / Mimari
s = slide("ATK verisinden Fuji kontrolüne.", "Piyasa verisi ATK CLI’dan TR sitesi ve JSON çıktısıyla geliyor; logger yirmi paritenin epizodik hafızasını altmış saniye aralıkla tutuyor. Deterministik metrikler kaymayı ve derinliği hesaplıyor, geçmiş medyanla piyasa rejimini karşılaştırıyor, LLM de bu araçların kullanımını planlıyor. Sonuç EIP-191 ile imzalanıp Fuji’de denetleniyor; piyasa verisi için yüz yirmi saniyelik eşik ve kontrattaki üç yüz saniyelik imza penceresi ayrı kontroller.", 40, light=True)
labels = [("ATK", "Canlı tahta"), ("Metrikler", "İcra maliyeti"), ("Planlayıcı", "LLM araç seçimi"), ("İmza + Fuji", "Devre kesici")]
for i, (head, body) in enumerate(labels):
    x = .85 + 3.05*i
    line(s, x, 2.52, x+2.5, 2.52, PAPER_MUTED, 1)
    text(s, x, 2.84, 2.56, .7, head, 27, INK, SERIF)
    text(s, x, 3.76, 2.55, .85, body, 22, PAPER_MUTED)
    if i<3:
        line(s, x+2.55, 3.24, x+2.91, 3.24, PAPER_MUTED, 1.5, True)
text(s, .85, 5.18, 5.3, .75, "20 parite", 36, INK, MONO)
text(s, .88, 6.03, 5.3, .5, "epizodik hafıza", 22, PAPER_MUTED)
text(s, 7.06, 5.18, 5.3, .75, "60 sn", 36, INK, MONO)
text(s, 7.09, 6.03, 5.3, .5, "kayıt aralığı", 22, PAPER_MUTED)

# 5 / Agentic yaklaşım
s = slide("LLM planlar. Sayılar deterministik.", "Agentic taraf, hangi kanıta hangi araçla bakılacağını seçen LLM planlayıcı; scan_market, get_candles, get_orderbook, compute_metric ve argümansız sign_score araçlarını kullanıyor. Sayılar deterministik hesaplanıyor, karar birleştirmede yalnızca daha muhafazakâr öneri kazanıyor ve sağlayıcılar çalışmazsa deterministik yedek plan devreye giriyor. Dürüst sınır şu: kontrat LLM karar etiketini değil imzalı sayıları denetler, dolayısıyla entegratör uygulama kararını da uygulamalı; serbest kullanıcı metni LLM’e gönderilmez.", 40)
for x, head, detail in [( .86,"LLM", "Aracı seçer.\nKanıtı yorumlar."), (5.05,"Metrikler", "Skoru hesaplar.\nGüvenli tutarı bulur."), (9.12,"Kontrat", "İmzayı doğrular.\nSınırı uygular.")]:
    line(s, x, 2.3, min(x+3.3,12.45), 2.3, RULE)
    text(s, x, 2.72, 3.35, .75, head, 36, SAGE, SERIF)
    text(s, x, 3.7, 3.3, 1.3, detail, 25, FG)
text(s, .9, 5.62, 11.7, .5, "Anthropic   /   Gemini   /   deterministik yedek", 23, SAGE)
text(s, .9, 6.37, 11.5, .35, "Sınır: LLM karar etiketi zincirde imzalanmıyor.", 17, MUTED)

# 6 / Demo
s = slide("Üç sahnede çalışan sistem.", "Şimdi yerel uygulamada önce bin USDT niyetini açıp karar izini, güvenli tutarı ve imzayı gösteriyorum; tahta uygunsa sonuç onay, canlı veri değiştiğinde sonuç da değişebilir. Sonra on bin USDT niyetinde güvenli tutar aşılırsa ExceedsSafeSize ve N parçaya bölme önerisini karşılaştırıyorum; N’yi önceden sabitlemiyorum, agent ve kontrat aynı yukarı yuvarlama hesabını kullanıyor. Üçüncü sahne açıkça veri kesintisi simülasyonu: oracle imza üretmiyor ve trader’ın kendi imzası BadSigner alıyor; varsayılan yerel kontrat cevabı eth_call ön izlemesi, gerçek işlem kanıtları ayrı ve internet sürümü yalnızca kayıtlı örnekleri gösteriyor.", 85, light=True)
for x, num, amount, decision, color, detail in [
    (.87,"01","1.000 USDT","ONAY",GREEN,"Tahta uygunsa\nimzalı kontrol"),
    (5.06,"02","10.000 USDT","BÖL",AMBER,"ExceedsSafeSize\nN parçaya böl"),
    (9.07,"03","Veri yok","İMZA YOK",RED,"SİMÜLASYON\nBadSigner")]:
    text(s, x, 2.05, 3.3, .4, num, 18, PAPER_MUTED, MONO)
    line(s, x, 2.64, min(x+3.25,12.42), 2.64, "B7C5BA")
    text(s, x, 2.95, 3.4, .75, amount, 29, INK)
    text(s, x, 4.02, 3.4, .66, decision, 27, color, bold=True)
    text(s, x, 4.98, 3.4, 1.05, detail, 22, INK)
text(s, .9, 6.43, 11.6, .38, "Yerelde canlı demo. İnternette kayıtlı simülasyon.", 20, PAPER_MUTED, link=DEMO)

# 7 / Kötü senaryo
s = slide("Veri bozulursa imza yok.", "Kötü senaryoda önce veri kesintisini gösteriyorum: imza yok, yetkisiz imza BadSigner alıyor; replay’de ise kontrat önce tazeliğe baktığı için eski kullanılmış onay StaleScore, geçerlilik penceresindeki kullanılmış onay AlreadyUsed alıyor ve iki yol da gerçek tx ile kaydedildi. Oracle anahtarı çalınırsa owner pause ve setOracle ile müdahale edebiliyor, ama tek oracle güveni devam ediyor. LLM için yapısal veri ve deterministik sayılar, yanlış alarm için parametrik eşik ve nominal Bonferroni düzeltmesi var; gerçek yanlış alarm oranını ölçtüğümü söylemiyorum ve acil çıkış yolu yalnızca kavram.", 45)
entries = [
    ("Veri kesintisi", "İmza üretilmez"),
    ("Yetkisiz imza", "BadSigner"),
    ("Eski onay / tekrar", "StaleScore / AlreadyUsed"),
    ("Oracle anahtarı", "pause + setOracle"),
    ("LLM / yanlış alarm", "Deterministik skor + parametrik eşik")]
for i,(a,b) in enumerate(entries):
    y=2.02+i*.79
    line(s,.87,y+.61,12.43,y+.61,RULE,.7)
    text(s,.9,y,4.4,.53,a,24,MUTED)
    text(s,5.42,y,7.0,.56,b,23,FG,MONO if i in (1,2,3) else SANS)
text(s,.9,6.43,11.6,.35,"Tazelik kontrolü replay kontrolünden önce gelir.",19,SAGE)

# 8 / Kanıt
s = slide("Zincirde doğrulandı.", "Ana kanıt listesinde yedi gerçek işlem var: IntentApproved kayıtları ve demo kapsamındaki ExceedsSafeSize, BadSigner, AlreadyUsed ve StaleScore kontrolleri; taze replay için bir USDT hazırlık işlemi ayrıca kullanıldı. Tablodaki işlem kimlikleri tıklanabilir, Python testlerinin güncel sayısı seksen yedi ve Solidity için on dört test kaydı var; burada Python koşuları ortam dosyası okunmadan, açık kontrat adresiyle doğrulandı, Solidity sayısı STATUS kaydından geliyor. Yaklaşık iki saniye zincir onayı ve koşu başına yaklaşık 0,001 USD LLM maliyeti proje kayıtlarına dayanıyor, bütün uygulamanın iki saniyede tamamlandığı anlamına gelmiyor; tek oracle ve testnet sınırları devam ediyor, risk ağırlıkları tasarım tercihi.", 45, light=True)
text(s,.86,1.86,7.7,.55,"7 gerçek tx / ana kanıt listesi",23,PAPER_MUTED)
for i,(label,tx) in enumerate(proofs):
    y=2.58+i*.48
    title=label.split(" (",1)[0]
    if i in (0,1):title += " / ilk" if i==0 else " / agent"
    if i in (2,3):title += " / ilk" if i==2 else " / agent"
    text(s,.88,y,4.13,.39,title,20,INK)
    text(s,5.08,y,3.18,.39,tx[:10]+"…"+tx[-6:],20,PAPER_MUTED,MONO,link=EXPLORER+tx)
    line(s,.88,y+.43,8.2,y+.43,"B7C5BA",.5)
line(s,8.68,2.0,8.68,6.3,"B7C5BA")
text(s,9.2,2.15,3.05,.85,"87/87",43,INK,MONO)
text(s,9.25,3.1,3.03,.45,"Python testi",22,PAPER_MUTED)
text(s,9.2,4.05,3.0,.85,"~2 sn",43,INK,MONO)
text(s,9.25,5.0,3.02,.65,"zincir onayı",22,PAPER_MUTED)
text(s,9.25,6.05,3.0,.4,"LLM ~0,001 USD/koşu",20,PAPER_MUTED)
text(s,.9,6.21,7.6,.36,"14/14 Solidity testi (STATUS kaydı)",18,PAPER_MUTED)
text(s,.9,6.64,11.6,.23,"Fuji kontratı: "+CONTRACT,13,PAPER_MUTED,MONO,link="https://explorer-test.avax.network/c-chain/address/"+CONTRACT)

# 9 / Ürün ve gelecek
s = slide("Bugün çalışan. Sonra geliştirilecek.", "Bugün yerelde çalışan canlı ATK ve imzalı skor akışı, Fuji kontratı ve paylaşılabilir bir internet simülasyonu var; internet sürümü kayıtlı örnekler kullanıyor, yeni veri veya imza üretmiyor. Hedef müşteri ATK ile agent geliştiren ekipler ve bireyler, tasarlanan iş modeli temel kontrolü ücretsiz tutup gerçek zamanlı API erişimini ücretlendirmek ama fiyat ve ödeme isteği henüz doğrulanmadı. Çoklu oracle, eşik imza, EIP-712, Almgren–Chriss icra planı, OKX.ai listeleme ve çağrı başına ödeme sonraki adımlar; bunları bugün kodlanmış özellik gibi sunmuyorum.",45)
text(s,.87,2.0,5.5,.55,"Bugün",30,SAGE,SERIF)
text(s,.87,2.98,5.55,1.8,"Yerelde canlı demo\nFuji’de devre kesici\nİnternette simülasyon",26,FG)
line(s,6.55,2.05,6.55,5.0,RULE)
text(s,7.12,2.0,5.3,.55,"Sonraki adım",30,SAGE,SERIF)
text(s,7.12,2.98,5.3,1.8,"Çoklu oracle + EIP-712\nİcra planı\nOKX.ai servis yayını",26,FG)
text(s,7.14,5.05,5.2,.35,"Henüz kodlanmadı.",18,MUTED)
line(s,.87,5.75,12.43,5.75,RULE)
text(s,.9,6.03,11.6,.6,"İş modeli tasarımı: temel ücretsiz, gerçek zamanlı API ücretli.",22,MUTED)

# 10 / Kapanış
s = slide("Teşekkürler.", "Teşekkür ederim, Arf’ın amacı her emirden önce tahtanın taşıyabileceği tutarı görünür ve doğrulanabilir kılmak. GitHub bağlantısında kaynak kod ve gerçek işlem kanıtları, internet bağlantısında ise kayıtlı verilerle çalışan etkileşimli simülasyon var. Canlı sistemi yerelde göstermeye ve sorularınızı cevaplamaya hazırım.",25,light=True,heading=False)
text(s,.8,1.15,11.6,1.35,"Teşekkürler.",74,INK,SERIF)
text(s,.89,2.93,11.5,.85,"Arf",40,INK,SERIF)
text(s,.9,3.8,11.4,.67,"İşlem öncesi icra ve likidite kalkanı.",28,PAPER_MUTED)
line(s,.9,4.85,12.4,4.85,"B7C5BA")
text(s,.92,5.23,10.8,.6,"github.com/akifaybek/arf",29,INK,link=GITHUB)
text(s,.94,6.07,11.0,.48,"Web demosu: kayıtlı verilerle etkileşimli simülasyon",21,PAPER_MUTED,link=DEMO)

out=ROOT/'Arf.pptx'
prs.save(out)
check=Presentation(out)
assert len(check.slides)==10
assert abs(check.slide_width/check.slide_height-16/9)<.001
all_text=[]
for s,record in zip(check.slides,records):
    visible=[]
    for shape in s.shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=check.slide_width+100
        assert shape.top+shape.height<=check.slide_height+100
        if shape.has_text_frame:visible.append(shape.text)
    assert check.slides[record['slide']-1].notes_slide.notes_text_frame.text==record['notes']
    record['visible_words']=len(' '.join(visible).split())
    all_text.extend(visible)
    print(f"{record['slide']:02} {record['title']} | {record['visible_words']} görünür kelime | not: {record['note_characters']} karakter")
joined=' '.join(all_text+[r['notes'] for r in records])
assert not re.search(r'sk-ant|AIza|PRIVATE_KEY|manipülasyon|wash trading|öğrenme',joined,re.I)
assert sum(r['target_seconds'] for r in records)==420
summary={'slides':10,'aspect':'16:9','palette':{'petrol':BG,'sage':SAGE,'paper':PAPER},
         'target_seconds_including_demo':420,'notes_words':sum(r['note_words'] for r in records),
         'slides_detail':records,'onchain_proofs':[{'event':a,'tx':b} for a,b in proofs],
         'preparation_tx':'0x0b4cfc04d3c91d5fe0fe09667d66eb07ef1bc630c8cec6b13ab47fe07ad23530',
         'github':GITHUB,'web_demo':DEMO}
(ROOT/'validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
(ROOT/'speaker_notes.md').write_text('# Arf konuşmacı notları\n\n10 slayt. Hedef 7 dakika, canlı demo dahil. Süre prova ile doğrulanacak.\n\n'+'\n\n'.join(f"## {r['slide']}. {r['title']} — {r['target_seconds']} sn\n\n{r['notes']}" for r in records)+'\n')
print(f'Üretildi: {out}; notlar {summary["notes_words"]} kelime, prova hedefi 420 sn.')
